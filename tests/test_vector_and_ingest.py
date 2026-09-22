import asyncio
import os
import shutil
import tempfile
import unittest
from pathlib import Path

import docx
import openpyxl
from pptx import Presentation

from bot.services.db_service import DatabaseService
from bot.services.ingest import DocumentIngestionService
from bot.services.vector_db import VectorDBService


class TestVectorAndIngest(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.chroma_dir = Path(self.temp_dir) / "chroma"
        self.db_path = Path(self.temp_dir) / "test.db"
        self.db_service = DatabaseService(self.db_path)
        await self.db_service.init_db()

        self.vector_db = VectorDBService(
            persist_dir=self.chroma_dir,
            model_name="sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"
        )
        await self.vector_db.initialize()
        self.ingest_service = DocumentIngestionService(default_chunk_size=100, default_overlap=20)

    async def asyncTearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    async def test_text_ingestion_and_user_isolation(self):
        user_1 = 1001
        user_2 = 2002

        # Создаем документ для пользователя 1
        doc_path = Path(self.temp_dir) / "sample_policy.txt"
        sample_text = (
            "Политика безопасности SchoolX регламентирует правила доступа сотрудников к корпоративным данным. "
            "Каждый сотрудник обязан использовать двухфакторную аутентификацию (2FA) для входа в рабочий кабинет. "
            "Секретный пароль пользователя 1: Секрет_Пользователя_1001."
        )
        doc_path.write_text(sample_text, encoding="utf-8")

        # Индексация для user_1
        chunks_added = await self.ingest_service.ingest_file(
            user_id=user_1,
            file_path=doc_path,
            original_filename="sample_policy.txt",
            vector_db=self.vector_db
        )
        self.assertGreater(chunks_added, 0)

        # Проверка: у user_1 есть чанки
        total_chunks_1 = await self.vector_db.count(user_1)
        self.assertEqual(total_chunks_1, chunks_added)

        # ПРОВЕРКА СТРОГОЙ ИЗОЛЯЦИИ: у user_2 в базе должно быть 0 чанков!
        total_chunks_2 = await self.vector_db.count(user_2)
        self.assertEqual(total_chunks_2, 0)

        # Поиск user_1 должен найти документ
        results_1 = await self.vector_db.similarity_search(user_id=user_1, query="Секретный пароль", k=2)
        self.assertGreater(len(results_1), 0)
        self.assertIn("Секрет_Пользователя_1001", results_1[0]["content"])

        # Поиск user_2 по тому же запросу должен вернуть ПУСТОЙ результат!
        results_2 = await self.vector_db.similarity_search(user_id=user_2, query="Секретный пароль", k=2)
        self.assertEqual(len(results_2), 0)

    async def test_delete_source_and_clear_all(self):
        user_id = 5555
        doc1 = Path(self.temp_dir) / "doc1.txt"
        doc1.write_text("Первый личный тестовый документ.", encoding="utf-8")
        doc2 = Path(self.temp_dir) / "doc2.txt"
        doc2.write_text("Второй документ с другой информацией.", encoding="utf-8")

        await self.ingest_service.ingest_file(user_id, doc1, "doc1.txt", self.vector_db)
        await self.ingest_service.ingest_file(user_id, doc2, "doc2.txt", self.vector_db)

        stats = await self.vector_db.get_sources_stats(user_id)
        self.assertEqual(len(stats), 2)

        # Удаление первого документа
        deleted = await self.vector_db.delete_source(user_id, "doc1.txt")
        self.assertGreater(deleted, 0)

        sources_after_del = await self.vector_db.get_unique_sources(user_id)
        self.assertNotIn("doc1.txt", sources_after_del)
        self.assertIn("doc2.txt", sources_after_del)

        # Очистка личной базы
        cleared = await self.vector_db.clear_all(user_id)
        self.assertGreater(cleared, 0)
        self.assertEqual(await self.vector_db.count(user_id), 0)

    async def test_ms_office_docx_ingestion_and_vectorization(self):
        user_id = 9001
        doc = docx.Document()
        doc.add_heading("Регламент безопасности SchoolX", level=1)
        doc.add_paragraph("Основное правило: двухфакторная аутентификация через приложение обязательна для всех сотрудников.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Сервис"
        table.cell(0, 1).text = "Порт"
        table.cell(1, 0).text = "LM Studio"
        table.cell(1, 1).text = "1234"

        docx_path = Path(self.temp_dir) / "security.docx"
        doc.save(docx_path)

        # Индексация с одновременной записью в SQLite БД
        added = await self.ingest_service.ingest_file(
            user_id=user_id,
            file_path=docx_path,
            original_filename="security.docx",
            vector_db=self.vector_db,
            db_service=self.db_service
        )
        self.assertGreater(added, 0)

        # Проверка ChromaDB
        self.assertEqual(await self.vector_db.count(user_id), added)

        # Проверка SQLite БД
        db_docs = await self.db_service.get_user_documents(user_id)
        self.assertEqual(len(db_docs), 1)
        self.assertEqual(db_docs[0]["filename"], "security.docx")
        self.assertEqual(db_docs[0]["file_type"], ".docx")
        self.assertEqual(db_docs[0]["chunks_count"], added)

        # Проверка векторного поиска по параграфу
        search_res1 = await self.vector_db.similarity_search(user_id=user_id, query="двухфакторная аутентификация", k=2)
        self.assertGreater(len(search_res1), 0)
        self.assertIn("двухфакторная аутентификация", search_res1[0]["content"])

        # Проверка векторного поиска по таблице
        search_res2 = await self.vector_db.similarity_search(user_id=user_id, query="порт сервиса LM Studio", k=2)
        self.assertGreater(len(search_res2), 0)
        self.assertIn("1234", search_res2[0]["content"])

    async def test_ms_office_xlsx_ingestion_and_vectorization(self):
        user_id = 9002
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Сотрудники"
        ws.append(["ФИО", "Специализация", "Телефон"])
        ws.append(["Дмитрий Волков", "Архитектор ИИ систем", "+7-999-123-45-67"])
        xlsx_path = Path(self.temp_dir) / "team.xlsx"
        wb.save(xlsx_path)

        added = await self.ingest_service.ingest_file(
            user_id=user_id,
            file_path=xlsx_path,
            original_filename="team.xlsx",
            vector_db=self.vector_db,
            db_service=self.db_service
        )
        self.assertGreater(added, 0)

        # Проверка SQLite БД
        db_docs = await self.db_service.get_user_documents(user_id)
        self.assertEqual(len(db_docs), 1)
        self.assertEqual(db_docs[0]["filename"], "team.xlsx")
        self.assertEqual(db_docs[0]["file_type"], ".xlsx")

        # Векторный поиск по содержимому таблицы Excel
        search_res = await self.vector_db.similarity_search(user_id=user_id, query="Архитектор ИИ систем Дмитрий", k=2)
        self.assertGreater(len(search_res), 0)
        self.assertIn("Дмитрий Волков", search_res[0]["content"])

    async def test_ms_office_pptx_ingestion_and_vectorization(self):
        user_id = 9003
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Курс SchoolX по искусственному интеллекту"
        slide.placeholders[1].text = "Тема 1: Нейронные сети и векторные базы данных"
        pptx_path = Path(self.temp_dir) / "ai_course.pptx"
        prs.save(pptx_path)

        added = await self.ingest_service.ingest_file(
            user_id=user_id,
            file_path=pptx_path,
            original_filename="ai_course.pptx",
            vector_db=self.vector_db,
            db_service=self.db_service
        )
        self.assertGreater(added, 0)

        # Проверка SQLite БД
        db_docs = await self.db_service.get_user_documents(user_id)
        self.assertEqual(len(db_docs), 1)
        self.assertEqual(db_docs[0]["filename"], "ai_course.pptx")

        # Векторный поиск по содержимому слайда PowerPoint
        search_res = await self.vector_db.similarity_search(user_id=user_id, query="векторные базы данных", k=2)
        self.assertGreater(len(search_res), 0)
        self.assertIn("векторные базы данных", search_res[0]["content"])

    async def test_odt_ingestion_and_vectorization(self):
        user_id = 9004
        import zipfile
        content_xml = (
            b'<?xml version="1.0" encoding="UTF-8"?>'
            b'<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
            b'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
            b'<office:body><office:text>'
            b'<text:h text:outline-level="1">\xd0\xa1\xd0\xbe\xd0\xb3\xd0\xbb\xd0\xb0\xd1\x88\xd0\xb5\xd0\xbd\xd0\xb8\xd0\xb5 \xd0\xbe \xd0\xba\xd0\xbe\xd0\xbd\xd1\x84\xd0\xb8\xd0\xb4\xd0\xb5\xd0\xbd\xd1\x86\xd0\xb8\xd0\xb0\xd0\xbb\xd1\x8c\xd0\xbd\xd0\xbe\xd1\x81\xd1\x82\xd0\xb8 SchoolX</text:h>'
            b'<text:p>\xd0\xa1\xd1\x82\xd1\x80\xd0\xbe\xd0\xb3\xd0\xbe \xd0\xb7\xd0\xb0\xd0\xbf\xd1\x80\xd0\xb5\xd1\x89\xd0\xb0\xd0\xb5\xd1\x82\xd1\x81\xd1\x8f \xd0\xbf\xd0\xb5\xd1\x80\xd0\xb5\xd0\xb4\xd0\xb0\xd0\xb2\xd0\xb0\xd1\x82\xd1\x8c \xd1\x81\xd0\xb5\xd0\xba\xd1\x80\xd0\xb5\xd1\x82\xd0\xbd\xd1\x8b\xd0\xb5 API-\xd0\xba\xd0\xbb\xd1\x8e\xd1\x87\xd0\xb8 \xd1\x82\xd1\x80\xd0\xb5\xd1\x82\xd1\x8c\xd0\xb8\xd0\xbc \xd0\xbb\xd0\xb8\xd1\x86\xd0\xb0\xd0\xbc.</text:p>'
            b'</office:text></office:body></office:document-content>'
        )
        odt_path = Path(self.temp_dir) / "nda.odt"
        with zipfile.ZipFile(odt_path, "w") as zf:
            zf.writestr("mimetype", "application/vnd.oasis.opendocument.text")
            zf.writestr("content.xml", content_xml)

        added = await self.ingest_service.ingest_file(
            user_id=user_id,
            file_path=odt_path,
            original_filename="nda.odt",
            vector_db=self.vector_db,
            db_service=self.db_service
        )
        self.assertGreater(added, 0)

        # Проверка SQLite БД
        db_docs = await self.db_service.get_user_documents(user_id)
        self.assertEqual(len(db_docs), 1)
        self.assertEqual(db_docs[0]["filename"], "nda.odt")
        self.assertEqual(db_docs[0]["file_type"], ".odt")

        # Векторный поиск по ODT документу
        search_res = await self.vector_db.similarity_search(user_id=user_id, query="передача секретных API ключей", k=2)
        self.assertGreater(len(search_res), 0)
        self.assertIn("API-ключи третьим лицам", search_res[0]["content"])


if __name__ == "__main__":
    unittest.main()
